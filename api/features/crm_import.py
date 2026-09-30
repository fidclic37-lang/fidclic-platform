"""
CRM › Import — businesses from a file, without any language model.

Reads every usual file (decided with the user: no LLM) — CSV, TSV, TXT,
Markdown, PDF, Excel (XLSX, XLSM, XLS), LibreOffice (ODS, ODT), Word (DOCX),
JSON, HTML, vCard (VCF), RTF, and any other text file. Two ways a file
carries businesses, both handled:

    A table      a spreadsheet, a CSV (« , » « ; » tab or « | »), a Markdown
                 or HTML table, a Word/LibreOffice table — or a table inside a
                 PDF, rebuilt from where each word sits on the page (a cell
                 that wraps over several lines stays one cell, a table that
                 goes on to the next page keeps its columns). Columns are
                 recognised by their header, in French or English (SYNONYMS),
                 wherever the header row is (titles above it are skipped),
                 and a file may hold several tables.
    Free text    one business per block (blank line, Markdown heading or list
                 item between them). In each block: « Clé : valeur » lines,
                 then patterns — e-mail, French phone, SIRET, postal code +
                 city, website, Instagram/Facebook links, and a sector read
                 from known words. The first line that is none of these is
                 the business's name.

The same business listed twice in the file is kept once (merged).

parse(filename, data) → [row] with the fields below; nothing is written. The
CRM decides what each row becomes (features/crm.py).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from typing import Dict, List, Optional

FIELDS = ("name", "sector", "city", "postal_code", "address", "contact_name", "contact_role", "phone",
          "email", "website", "instagram", "facebook", "siret", "status", "source", "notes")

SYNONYMS = {
    "name": ["nom du commerce", "commerce", "nom commercial", "entreprise", "societe", "raison sociale", "etablissement",
             "enseigne", "business", "company", "shop", "store", "nom", "name", "restaurant", "restaurants", "boutique",
             "magasin", "denomination", "nom de l'etablissement", "nom etablissement", "client", "prospect"],
    "sector": ["secteur", "activite", "categorie", "type", "sector", "industry", "category", "segment", "sous-segment",
               "sous segment", "specialite", "cuisine", "type de commerce"],
    "city": ["ville", "commune", "city", "localite"],
    "postal_code": ["code postal", "cp", "postal code", "zip", "postcode"],
    "address": ["adresse", "rue", "address", "street"],
    "contact_name": ["contact", "gerant", "nom du contact", "interlocuteur", "responsable", "proprietaire",
                     "contact name", "owner", "prenom nom"],
    "contact_role": ["fonction", "poste", "role", "title"],
    "phone": ["telephone", "tel", "portable", "mobile", "phone", "numero"],
    "email": ["email", "e-mail", "mail", "courriel"],
    "website": ["site", "site web", "website", "url", "web"],
    "instagram": ["instagram", "insta"],
    "facebook": ["facebook", "fb"],
    "siret": ["siret", "siren"],
    "status": ["statut", "status", "etape", "stage"],
    "source": ["source", "origine", "provenance"],
    "notes": ["notes", "note", "commentaire", "commentaires", "remarques", "comments"],
}

SECTOR_WORDS = [
    ("Restauration / Fast-food", ["fast-food", "fast food", "restaurant", "restauration", "burger", "kebab", "tacos", "sushi",
                                  "pizzeria", "pizza", "brasserie", "snack", "traiteur", "bistrot", "creperie"]),
    ("Café / Salon de thé", ["cafe", "salon de the", "coffee", "coffee shop", "torrefacteur"]),
    ("Barbershop", ["barber", "barbershop", "barbier"]),
    ("Coiffure / Beauté", ["coiffure", "coiffeur", "beaute", "esthetique", "onglerie", "institut", "salon de beaute"]),
    ("Boulangerie", ["boulangerie", "patisserie", "viennoiserie", "boulanger"]),
    ("Bien-être", ["spa", "massage", "bien-etre", "bien etre", "yoga", "sauna"]),
    ("Mode", ["mode", "pret-a-porter", "vetements", "boutique", "chaussures", "fashion"]),
    ("Salle de sport", ["salle de sport", "fitness", "gym", "crossfit", "musculation"]),
    ("Optique", ["optique", "opticien", "lunettes"]),
    ("Bar", ["bar", "pub", "cave a vin", "bar a vin"]),
    ("Épicerie / Commerce de proximité", ["epicerie", "superette", "primeur", "fromagerie", "boucherie", "caviste", "tabac"]),
]

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?:(?:\+|00)33[\s.\-]?|0)[1-9](?:[\s.\-]?\d{2}){4}")
SIRET_RE = re.compile(r"\b\d{3}\s?\d{3}\s?\d{3}(?:\s?\d{5})?\b")
CP_CITY_RE = re.compile(r"\b(\d{5})\s+([A-Za-zÀ-ÿ'’\- ]{2,40}?)(?=$|,|\n|\s{2}|\s-\s|\s\()")
URL_RE = re.compile(r"(?:https?://|www\.)[^\s,;)]+", re.I)
KV_RE = re.compile(r"^\s*[-*•]?\s*\**([A-Za-zÀ-ÿ .'’/-]{2,30}?)\**\s*[:=]\s*(.+)$")


def _fold(s: str) -> str:
    s = unicodedata.normalize("NFD", (s or "").lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn").strip()


def field_for(header: str) -> Optional[str]:
    h = re.sub(r"[_*#()\[\]]", " ", _fold(header))
    h = re.sub(r"\s+", " ", h).strip()
    if not h:
        return None
    for field, words in SYNONYMS.items():
        if h in words:
            return field
    for field, words in SYNONYMS.items():
        if any(len(w) > 3 and w in h for w in words):
            return field
    return None


def sector_of(text: str) -> str:
    t = f" {_fold(text)} "
    for label, words in SECTOR_WORDS:
        if any(re.search(rf"(?<![a-z]){re.escape(w)}(?![a-z])", t) for w in words):
            return label
    return ""


EMPTY = {"—", "–", "-", "--", "n/a", "na", "nc", "n.c.", "néant", "neant", "aucun", "aucune", "/", "?", "none", "null", "nan"}


def _cell(v) -> str:
    """A spreadsheet or table value as text: numbers without « .0 »."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = re.sub(r"\s+", " ", str(v)).strip()
    return "" if t.lower() in EMPTY else t


def _clean_phone(p: str) -> str:
    d = re.sub(r"[^\d+]", "", p or "")
    if re.fullmatch(r"[1-9]\d{8}", d):
        d = "0" + d                                   # a spreadsheet dropped the leading 0
    if d.startswith("0033"):
        d = "+33" + d[4:]
    if d.startswith("+33") and len(d) == 12:
        d = "0" + d[3:]
    if len(d) == 10 and d.startswith("0"):
        return " ".join(d[i:i + 2] for i in range(0, 10, 2))
    return (p or "").strip()


def _finish(row: Dict[str, str]) -> Dict[str, str]:
    out = {k: _cell(row.get(k)) for k in FIELDS}
    # « Tours » alone in the address column is a city, not an address.
    if out["address"] and not out["city"] and not re.search(r"\d", out["address"]) \
            and re.fullmatch(r"[A-ZÀ-Ý][\wÀ-ÿ'’-]+", out["address"]):
        out["city"], out["address"] = out["address"], ""
    if out["email"]:
        m = EMAIL_RE.search(out["email"])
        out["email"] = m.group(0).lower() if m else ""
    out["phone"] = _clean_phone(out["phone"])
    out["siret"] = re.sub(r"\D", "", out["siret"])
    if out["postal_code"]:
        m = re.search(r"\d{5}", out["postal_code"])
        out["postal_code"] = m.group(0) if m else ""
    if not out["postal_code"] or not out["city"]:
        m = CP_CITY_RE.search(f"{out['address']} {out['city']}")
        if m:
            out["postal_code"] = out["postal_code"] or m.group(1)
            out["city"] = out["city"] or m.group(2).strip().title()
    if not out["sector"]:
        out["sector"] = sector_of(f"{out['name']} {out['notes']}")
    for k in ("instagram", "facebook", "website"):
        if out[k] and not out[k].startswith("http") and "." in out[k]:
            out[k] = "https://" + out[k].lstrip("/")
    return out


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------
def header_of(cells: List[str]) -> Optional[List[Optional[str]]]:
    """This row is a header when it names a business column (or phone /
    e-mail) plus at least one other, and carries no figures — a data row
    has a phone number or an address with digits."""
    if any(len(re.findall(r"\d", c or "")) >= 3 for c in cells):
        return None
    fields = [(None if (c or "").strip() in ("#", "n°", "no", "n") else field_for(c or "")) for c in cells]
    got = {f for f in fields if f}
    if len(got) >= 2 and ("name" in got or {"email", "phone"} & got):
        return fields
    return None


def grid_rows(grid: List[List[str]]) -> List[Dict[str, str]]:
    """Rows of a table (spreadsheet, CSV, Word…): the header can be anywhere,
    and a new header starts a new table."""
    rows, header = [], None
    for raw in grid:
        cells = [_cell(c) for c in raw]
        if not any(cells):
            continue
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue                                    # Markdown separator row
        h = header_of(cells)
        if h:
            header = h
            continue
        if not header:
            continue
        row: Dict[str, str] = {}
        for f, c in zip(header, cells):
            if f and c:
                row[f] = f"{row[f]} {c}".strip() if f in row else c
        if row.get("name") or row.get("email") or row.get("phone"):
            rows.append(_finish(row))
    return rows


def _table(lines: List[str], delim: str) -> Optional[List[Dict[str, str]]]:
    rows = grid_rows(list(csv.reader(lines, delimiter=delim)))
    return rows or None


def _sniff(text: str) -> Optional[List[Dict[str, str]]]:
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return []
    md = [ln.strip().strip("|") for ln in lines if ln.strip().startswith("|") and ln.count("|") >= 3]
    if len(md) >= 2:
        got = _table([re.sub(r"\s*\|\s*", "|", ln) for ln in md], "|")
        if got is not None:
            return got
    for delim in (";", "\t", ",", "|"):
        if sum(1 for ln in lines if ln.count(delim) >= 1) >= 2:   # titles above the table are fine
            got = _table(lines, delim)
            if got is not None:
                return got
    return None


# ---------------------------------------------------------------------------
# Free text
# ---------------------------------------------------------------------------
def _blocks(text: str) -> List[str]:
    """One block per business. With Markdown headings, each heading opens a
    business and everything under it belongs to it. Otherwise blank lines
    separate businesses — and a list whose every line is a whole business
    (« - Café X, Tours, 06 … ») is split line by line."""
    text = text.replace("\r", "")
    if re.search(r"^#{2,6}\s", text, re.M) or len(re.findall(r"^#\s", text, re.M)) > 1:
        parts = re.split(r"\n(?=#{1,6}\s)", "\n" + text)
    else:
        parts = re.split(r"\n\s*\n", text)
    out: List[str] = []
    for p in parts:
        lines = [ln for ln in p.strip().splitlines() if ln.strip()]
        items = [ln for ln in lines if re.match(r"\s*(?:\d+[.)]|[-*•])\s+", ln)]
        if len(items) >= 2 and len(items) == len(lines) and \
                all(EMAIL_RE.search(ln) or PHONE_RE.search(ln) for ln in items):
            out.extend(items)
        elif lines:
            out.append("\n".join(lines))
    return out


def _from_block(block: str) -> Optional[Dict[str, str]]:
    row: Dict[str, str] = {}
    leftovers: List[str] = []
    for raw in block.splitlines():
        line = raw.strip()
        if not line:
            continue
        m = KV_RE.match(line)
        f = field_for(m.group(1)) if m else None
        if m and f:
            row.setdefault(f, m.group(2).strip().strip("*"))
            continue
        leftovers.append(line)
    rest = "\n".join(leftovers)
    if "email" not in row and (m := EMAIL_RE.search(rest)):
        row["email"] = m.group(0)
    if "phone" not in row and (m := PHONE_RE.search(rest)):
        row["phone"] = m.group(0)
    for u in URL_RE.findall(rest):
        u = u.rstrip(".")
        key = "instagram" if "instagram." in u else "facebook" if "facebook." in u else "website"
        row.setdefault(key, u)
    if "siret" not in row and (m := re.search(r"siret\D{0,5}(\d[\d ]{12,16}\d)", rest, re.I)):
        row["siret"] = m.group(1)
    if "postal_code" not in row and (m := CP_CITY_RE.search(rest)):
        row["postal_code"] = m.group(1)
        row.setdefault("city", m.group(2).strip().title())
    if "name" not in row:
        for line in leftovers:
            cand = re.sub(r"^(?:#{1,6}\s*|\d+[.)]\s*|[-*•]\s*)", "", line).strip().strip("*_").strip()
            # The name is what comes before the first comma or dash —
            # « Kebab Istanbul, restauration rapide, 07 88 … ».
            head = re.split(r"\s[—–-]\s|,", cand)[0].strip()
            if not head or EMAIL_RE.search(head) or PHONE_RE.search(head) or URL_RE.search(head) \
                    or re.fullmatch(r"[\d\s,.-]+", head):
                continue
            row["name"] = head[:100]
            if head != cand and not row.get("contact_name"):
                # « Optique Vision — Camille B., 06 … »: a person's name after the dash.
                m = re.search(r"[—–-]\s*([A-ZÀ-Ý][a-zà-ÿ]+(?:\s[A-ZÀ-Ý][a-zà-ÿ]*\.?)+)", cand)
                if m:
                    row["contact_name"] = m.group(1)
            break
    if "address" not in row:
        for line in leftovers:
            if re.search(r"\b\d{1,4}\s*(?:bis|ter)?,?\s+(?:rue|avenue|av\.|boulevard|bd|place|quai|allée|chemin|route|impasse)\b", line, re.I):
                row["address"] = line.strip("-*• ")
                break
    if not row.get("sector"):
        row["sector"] = sector_of(block)
    if not (row.get("name") and (row.get("email") or row.get("phone") or row.get("city") or row.get("address"))):
        return None
    return _finish(row)


def parse_text(text: str) -> List[Dict[str, str]]:
    table = _sniff(text)
    if table is not None:
        return table
    return [r for r in (_from_block(b) for b in _blocks(text)) if r]


# ---------------------------------------------------------------------------
# PDF — tables rebuilt from where each word sits
# ---------------------------------------------------------------------------
def pdf_text(data: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "") for p in reader.pages)


def _segments(line: str) -> List[tuple]:
    """(column, text) of each cell on a laid-out line: cells are 2+ spaces apart."""
    return [(m.start(), m.group(0)) for m in re.finditer(r"\S+(?: \S+)*", line)]


def _join(a: str, b: str, field: Optional[str]) -> str:
    if not a:
        return b
    if field == "phone" or (re.fullmatch(r"[\d +().-]+", a) and re.fullmatch(r"[\d +().-]+", b)):
        return f"{a} {b}"
    if re.search(r"[A-Za-zÀ-ÿ]-$", a):               # « Bernard- » + « Palissy »
        return a + b
    return f"{a} {b}"


def pdf_tables(data: bytes) -> List[Dict[str, str]]:
    """Tables of a PDF: pypdf's layout mode keeps each word at its column, so
    a header row gives the columns, a row starts on its number (or its phone)
    and the lines below it go on filling the same cells. The columns are
    re-measured on every full row, so a table that continues on the next
    page — often shifted by a few characters — keeps its columns."""
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    out: List[Dict[str, str]] = []
    fields: Optional[List[Optional[str]]] = None       # the current table's columns
    xs: List[int] = []                                 # where they start
    row: Optional[List[str]] = None

    def flush():
        nonlocal row
        if row and fields:
            rec: Dict[str, str] = {}
            for f, v in zip(fields, row):
                if f and f != "#" and v:
                    rec[f] = f"{rec[f]} {v}".strip() if f in rec else v
            if rec.get("name") or rec.get("phone") or rec.get("email"):
                out.append(_finish(rec))
        row = None

    def col_of(x: int) -> int:
        left = [i for i, c in enumerate(xs) if c <= x + 2]
        return left[-1] if left else 0

    for page in reader.pages:
        try:
            lines = (page.extract_text(extraction_mode="layout") or "").splitlines()
        except Exception:
            lines = (page.extract_text() or "").splitlines()
        i = 0
        while i < len(lines):
            segs = _segments(lines[i])
            i += 1
            if not segs:
                continue
            texts = [t for _, t in segs]
            h = header_of(texts)
            if h:
                flush()
                # A header cell that wraps (« Sous- » / « segment »): join the line below.
                if i < len(lines):
                    nxt = _segments(lines[i])
                    if nxt and not any(re.search(r"\d", t) for _, t in nxt) and \
                            all(any(abs(x - sx) <= 3 for sx, _ in segs) for x, _ in nxt):
                        merged = [t + next((u for x, u in nxt if abs(x - sx) <= 3), "") if t.endswith("-")
                                  else t + (" " + next((u for x, u in nxt if abs(x - sx) <= 3), "")).rstrip()
                                  for sx, t in segs]
                        h2 = header_of(merged)
                        if h2:
                            h = h2
                            i += 1
                fields = [("#" if t.strip() in ("#", "N°", "n°", "No", "N") else f) for t, f in zip(texts, h)]
                xs = [x for x, _ in segs]
                continue
            if not fields:
                continue
            numbered = "#" in fields
            first_x, first_t = segs[0]
            starts = False
            if numbered:
                starts = bool(re.fullmatch(r"\d{1,4}\.?", first_t)) and col_of(first_x) <= fields.index("#")
            else:
                cells = {fields[col_of(x)]: t for x, t in segs}
                phone, email = cells.get("phone", ""), cells.get("email", "")
                starts = bool(re.match(r"(\+|0)\d", phone) or phone in EMPTY or "@" in email) or len(segs) >= max(2, len(fields) - 1)
            if starts:
                flush()
                if len(segs) == len(fields):
                    xs = [x for x, _ in segs]                  # re-measure the columns on a full row
                row = [""] * len(fields)
            elif row is None:
                # Not in a row: a title or a paragraph between tables ends the table.
                if len(segs) == 1 and first_x < (xs[1] if len(xs) > 1 else xs[0]) - 2 and len(first_t.split()) >= 1 \
                        and not re.search(r"\d", first_t):
                    fields = None
                continue
            for x, t in segs:
                c = col_of(x)
                row[c] = _join(row[c], "" if t in EMPTY else t, fields[c])
        # a row may go on at the top of the next page: keep it open
    flush()
    return out


# ---------------------------------------------------------------------------
# Spreadsheets and documents
# ---------------------------------------------------------------------------
def xlsx_grids(data: bytes) -> List[List[List[str]]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [[[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets]


def xls_grids(data: bytes) -> List[List[List[str]]]:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    return [[[_cell(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)] for sh in wb.sheets()]


def _xml(data: bytes, member: str):
    import zipfile
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return ET.fromstring(z.read(member))


def _text_of(el) -> str:
    return " ".join(t.strip() for t in el.itertext() if t and t.strip())


def odf_grids(data: bytes) -> tuple:
    """ODS / ODT: every table, plus the text outside tables."""
    T = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
    X = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    root = _xml(data, "content.xml")
    grids = []
    for tb in root.iter(f"{T}table"):
        grid = []
        for tr in tb.iter(f"{T}table-row"):
            cells = []
            for tc in tr:
                if tc.tag not in (f"{T}table-cell", f"{T}covered-table-cell"):
                    continue
                n = min(int(tc.get(f"{T}number-columns-repeated", "1") or 1), 60)
                cells += [_text_of(tc)] * n
            reps = min(int(tr.get(f"{T}number-rows-repeated", "1") or 1), 5)
            if any(cells):
                grid += [cells] * reps
        grids.append(grid)
    in_tables = {id(p) for tb in root.iter(f"{T}table") for p in tb.iter(f"{X}p")}
    text = "\n".join(_text_of(p) for p in root.iter(f"{X}p") if id(p) not in in_tables)
    return grids, text


def docx_grids(data: bytes) -> tuple:
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = _xml(data, "word/document.xml")
    grids = []
    in_tables = set()
    for tbl in root.iter(f"{W}tbl"):
        grid = []
        for tr in tbl.iter(f"{W}tr"):
            cells = []
            for tc in tr.findall(f"{W}tc"):
                cells.append(" ".join("".join(t.text or "" for t in p.iter(f"{W}t")) for p in tc.iter(f"{W}p")).strip())
            grid.append(cells)
        grids.append(grid)
        in_tables |= {id(p) for p in tbl.iter(f"{W}p")}
    paras = ["".join(t.text or "" for t in p.iter(f"{W}t")) for p in root.iter(f"{W}p") if id(p) not in in_tables]
    return grids, "\n".join(paras)


def html_grids(text: str) -> tuple:
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.grids, self.grid, self.row, self.cell, self.depth, self.other = [], None, None, None, 0, []

        def handle_starttag(self, tag, attrs):
            if tag == "table":
                self.depth += 1
                self.grid = []
            elif tag == "tr" and self.grid is not None:
                self.row = []
            elif tag in ("td", "th") and self.row is not None:
                self.cell = []
            elif tag in ("br", "p", "div", "li") and self.cell is None:
                self.other.append("\n")

        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.cell is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.cell = None
            elif tag == "tr" and self.row is not None:
                self.grid.append(self.row)
                self.row = None
            elif tag == "table" and self.grid is not None:
                self.grids.append(self.grid)
                self.grid = None
                self.depth -= 1

        def handle_data(self, d):
            if self.cell is not None:
                self.cell.append(d)
            elif self.grid is None:
                self.other.append(d)
    p = P()
    p.feed(text)
    return p.grids, "".join(p.other)


def json_rows(text: str) -> List[Dict[str, str]]:
    import json
    doc = json.loads(text)
    items = doc if isinstance(doc, list) else next((v for v in (doc or {}).values() if isinstance(v, list)), []) \
        if isinstance(doc, dict) else []
    rows = []
    for it in items:
        if not isinstance(it, dict):
            continue
        flat = {}
        for k, v in it.items():
            if isinstance(v, dict):
                flat.update({f"{k} {k2}": v2 for k2, v2 in v.items()})
            else:
                flat[k] = v
        row: Dict[str, str] = {}
        for k, v in flat.items():
            f = field_for(str(k))
            if f and _cell(v) and f not in row:
                row[f] = _cell(v)
        if row.get("name") or row.get("email") or row.get("phone"):
            rows.append(_finish(row))
    return rows


def vcf_rows(text: str) -> List[Dict[str, str]]:
    rows = []
    for card in re.split(r"(?im)^END:VCARD", text):
        if "BEGIN:VCARD" not in card.upper():
            continue
        get = lambda key: next((ln.split(":", 1)[1].strip() for ln in card.splitlines()   # noqa: E731
                                if ln.upper().split(";")[0].split(":")[0] == key and ":" in ln), "")
        adr = [x.strip() for x in get("ADR").split(";")] if get("ADR") else []
        row = {"name": get("ORG").replace(";", " ").strip() or get("FN"), "contact_name": get("FN") if get("ORG") else "",
               "contact_role": get("TITLE"), "phone": get("TEL"), "email": get("EMAIL"), "website": get("URL"),
               "address": adr[2] if len(adr) > 2 else "", "city": adr[3] if len(adr) > 3 else "",
               "postal_code": adr[5] if len(adr) > 5 else "", "notes": get("NOTE")}
        if row["name"] or row["phone"] or row["email"]:
            rows.append(_finish(row))
    return rows


def rtf_text(text: str) -> str:
    text = re.sub(r"\\par[d]?", "\n", text)
    text = re.sub(r"\\'([0-9a-f]{2})", lambda m: bytes.fromhex(m.group(1)).decode("cp1252", "ignore"), text)
    text = re.sub(r"\\[a-z]+-?\d* ?|[{}]", "", text)
    return text


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            t = data.decode(enc)
            if enc == "utf-16" and not data[:2] in (b"\xff\xfe", b"\xfe\xff"):
                continue
            return t
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------------------
# One business once
# ---------------------------------------------------------------------------
def _key(r: Dict[str, str]) -> str:
    return re.sub(r"[^a-z0-9]", "", _fold(r.get("name") or "")) or re.sub(r"\D", "", r.get("phone") or "") or (r.get("email") or "")


def dedupe(rows: List[Dict[str, str]]) -> tuple:
    """The same business listed twice in the file (same name, or no name and
    the same phone / e-mail) is kept once, its details merged."""
    seen: Dict[str, Dict[str, str]] = {}
    out, merged = [], 0
    for r in rows:
        k = _key(r)
        if k and k in seen:
            keep = seen[k]
            for f, v in r.items():
                if v and not keep.get(f):
                    keep[f] = v
            merged += 1
            continue
        if k:
            seen[k] = r
        out.append(r)
    return out, merged


UNSUPPORTED = (".doc", ".ppt", ".pptx", ".key", ".pages", ".numbers", ".zip", ".rar", ".7z", ".png", ".jpg", ".jpeg",
               ".gif", ".webp", ".heic", ".mp4", ".mov", ".mp3", ".exe", ".dmg")
ACCEPTED = ".csv, .tsv, .txt, .md, .pdf, .xlsx, .xlsm, .xls, .ods, .odt, .docx, .json, .html, .vcf, .rtf"


def parse_with_stats(filename: str, data: bytes) -> tuple:
    rows = parse_raw(filename, data)
    return dedupe(rows)


def parse(filename: str, data: bytes) -> List[Dict[str, str]]:
    return parse_with_stats(filename, data)[0]


def parse_raw(filename: str, data: bytes) -> List[Dict[str, str]]:
    name = (filename or "").lower()
    if name.endswith(UNSUPPORTED):
        raise ValueError(f"format non pris en charge — formats acceptés : {ACCEPTED}")
    if name.endswith(".pdf") or data[:4] == b"%PDF":
        rows = pdf_tables(data)
        return rows if rows else parse_text(pdf_text(data))
    if name.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")) or (data[:2] == b"PK" and b"xl/" in data[:2000]):
        return [r for g in xlsx_grids(data) for r in grid_rows(g)]
    if name.endswith((".xls", ".xlt")) or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return [r for g in xls_grids(data) for r in grid_rows(g)]
    if name.endswith((".ods", ".odt", ".ots")):
        grids, text = odf_grids(data)
        rows = [r for g in grids for r in grid_rows(g)]
        return rows or parse_text(text)
    if name.endswith(".docx") or (data[:2] == b"PK" and b"word/" in data[:2000]):
        grids, text = docx_grids(data)
        rows = [r for g in grids for r in grid_rows(g)]
        return rows or parse_text(text)
    text = _decode(data)
    if name.endswith(".json") or text.lstrip()[:1] in ("[", "{") and name.endswith((".json", ".txt", "")):
        try:
            return json_rows(text)
        except ValueError:
            pass
    if name.endswith((".vcf", ".vcard")) or "BEGIN:VCARD" in text[:2000].upper():
        return vcf_rows(text)
    if name.endswith((".html", ".htm", ".xhtml")) or re.search(r"<table[\s>]", text[:50000], re.I):
        grids, other = html_grids(text)
        rows = [r for g in grids for r in grid_rows(g)]
        return rows or parse_text(other)
    if name.endswith(".rtf") or text.startswith("{\\rtf"):
        text = rtf_text(text)
    return parse_text(text)
