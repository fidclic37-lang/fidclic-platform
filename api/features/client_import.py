"""
Clients › Importer — bring an existing client list in, and invite each client
to add their card (user, 2026-10-02).

THE FLOW
    1. The merchant drops a CSV or Excel file. `preview` reads it, finds the
       columns by their names (Nom, Prénom, Email, Téléphone, Date de
       naissance, Code postal, Ville — in French or English, any order) and
       says, line by line, what will happen: new card, already a client,
       twice in the file, or not usable (and why).
    2. `commit` re-checks every line on the server (the browser is never
       trusted), creates the cards within the plan's limits — active clients
       and new cards this month — and, if asked, emails each new client a
       link to their card.
    3. Clients who have not opened their card a week later are listed under
       « Relance » — prepared, never sent on its own: the merchant reads the
       list and decides (user, 2026-10-02).

CONSENT
    An imported client did not tick the join page's consent box, so the
    merchant must confirm they have their clients' agreement before the
    import is accepted. That confirmation is stored on each card as its
    consent record (source « import », who, when). Every email carries an
    unsubscribe link; an unsubscribed client is never emailed again.

QUOTA
    Invitation emails, relances included, count against the plan's
    `invite_emails_per_month` (Starter 200 / Pro 1 000 / Pro AI 3 000,
    editable in Super admin › Plan Limits). Only emails actually handed to
    Brevo count. Without BREVO_API_KEY the import still works; the page says
    emails are not switched on.

Mounts at:
    POST /api/owner/clients/import/preview           (multipart file)
    POST /api/owner/clients/import/commit
    GET  /api/owner/clients/invitations              quota, totals, relance count
    GET  /api/owner/clients/invitations/relance      who would get one
    POST /api/owner/clients/invitations/relance      send it to the chosen ones
    GET  /api/public/email/unsubscribe/{token}       one click, no login
"""
from __future__ import annotations

import csv
import html
import io
import re
import secrets
import unicodedata
import uuid
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional, Tuple

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse
from pydantic import BaseModel, Field

from auth import require_role, TokenData
from features import store_scope as S

router = APIRouter(tags=["clients"])
_db = None

IMPORTERS = ["business_owner", "administrateur"]
READERS = ["business_owner", "administrateur", "manager"]
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
RELANCE_AFTER_DAYS = 7

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")

# Header names, normalised (lower case, no accents, letters only), → field.
HEADERS = {
    "name": {"nomcomplet", "nometprenom", "prenometnom", "client", "fullname", "name", "nomprenom"},
    "first": {"prenom", "firstname", "first", "givenname"},
    "last": {"nom", "lastname", "last", "surname", "nomdefamille", "familyname"},
    "email": {"email", "mail", "courriel", "adresseemail", "adressemail", "emailaddress"},
    "phone": {"telephone", "tel", "phone", "mobile", "portable", "numero", "numerodetelephone", "phonenumber", "gsm"},
    "birthday": {"datedenaissance", "naissance", "anniversaire", "birthday", "dateofbirth", "birthdate", "dob"},
    "postal_code": {"codepostal", "cp", "postal", "postalcode", "zip", "zipcode"},
    "city": {"ville", "city", "commune", "localite"},
}
STATUS_LABELS = {"new": "Nouvelle carte", "existing": "Déjà client", "duplicate": "En double dans le fichier",
                 "invalid": "Ligne inutilisable"}


def init(db):
    global _db
    _db = db
    try:
        db.client_invitations.create_index([("tenant_id", 1), ("sent_at", -1)])
        db.client_invitations.create_index([("customer_id", 1), ("kind", 1)])
    except Exception:
        pass


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _aware(value: Any) -> Optional[datetime]:
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=timezone.utc)
    return None


# ---------------------------------------------------------------------------
# Reading the file
# ---------------------------------------------------------------------------
def _key(header: Any) -> str:
    s = unicodedata.normalize("NFKD", str(header or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z]", "", s)


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _table(name: str, raw: bytes) -> List[List[str]]:
    """The file as rows of text cells, header first."""
    lower = (name or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook
            ws = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).active
            return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
        except Exception:
            raise HTTPException(status_code=400, detail="Ce fichier Excel ne peut pas être lu. Enregistrez-le en .xlsx ou en .csv.")
    if lower.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Ancien format Excel (.xls) : enregistrez le fichier en .xlsx ou en .csv.")
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="Le fichier ne peut pas être lu.")
    sample = text[:4096]
    delim = max((";", ",", "\t"), key=lambda d: sample.split("\n", 1)[0].count(d))
    return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _columns(header: List[str]) -> Dict[str, int]:
    found: Dict[str, int] = {}
    for i, h in enumerate(header):
        k = _key(h)
        for field, names in HEADERS.items():
            if k in names and field not in found:
                found[field] = i
                break
    return found


def _phone(v: str) -> str:
    v = (v or "").strip()
    digits = re.sub(r"\D", "", v)
    if not digits:
        return ""
    if v.startswith("+"):
        return "+" + digits
    if len(digits) == 9 and not digits.startswith("0"):        # Excel ate the leading 0
        digits = "0" + digits
    return digits


def _birthday(v: str) -> Optional[str]:
    """MM-DD, as the join page stores it — or None."""
    v = (v or "").strip()
    if not v:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", v)                 # 1990-05-17
    if m:
        mm, dd = int(m.group(2)), int(m.group(3))
    else:
        m = re.match(r"^(\d{1,2})[/.\-](\d{1,2})(?:[/.\-](\d{2,4}))?$", v)   # 17/05/1990, 17/05
        if not m:
            return None
        dd, mm = int(m.group(1)), int(m.group(2))
    if not (1 <= mm <= 12 and 1 <= dd <= 31):
        return None
    return f"{mm:02d}-{dd:02d}"


def _mmdd(v: Optional[str]) -> Optional[str]:
    """A birthday sent back by the page: already MM-DD, or as typed in the file."""
    v = (v or "").strip()
    if re.match(r"^\d{2}-\d{2}$", v):
        mm, dd = int(v[:2]), int(v[3:])
        return v if 1 <= mm <= 12 and 1 <= dd <= 31 else None
    return _birthday(v)


def _row(cells: List[str], cols: Dict[str, int]) -> Dict[str, Any]:
    get = lambda f: cells[cols[f]].strip() if f in cols and cols[f] < len(cells) else ""
    name = get("name") or " ".join(x for x in (get("first"), get("last")) if x)
    return {"name": re.sub(r"\s+", " ", name)[:80], "email": get("email").lower()[:120],
            "phone": _phone(get("phone"))[:20], "birthday": _birthday(get("birthday")),
            "postal_code": re.sub(r"\s", "", get("postal_code"))[:10], "city": get("city")[:80]}


def _problem(r: Dict[str, Any]) -> Optional[str]:
    if not r["name"]:
        return "Nom manquant"
    if not r["email"] and not r["phone"]:
        return "Ni email ni téléphone"
    if r["email"] and not EMAIL_RE.match(r["email"]):
        return "Email invalide"
    if r["phone"] and not (8 <= len(r["phone"].lstrip("+")) <= 15):
        return "Téléphone invalide"
    return None


def _known(tid: str) -> Tuple[set, set]:
    """The emails and phone numbers this shop already has a card for."""
    emails, phones = set(), set()
    for c in _db.customers.find({"tenant_id": tid, "deleted_at": None}, {"_id": 0, "email": 1, "phone": 1}):
        if c.get("email"):
            emails.add(str(c["email"]).strip().lower())
        if c.get("phone"):
            phones.add(_phone(str(c["phone"])))
    return emails, phones


def classify(tid: str, rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Each row with its fate: new / existing / duplicate / invalid."""
    emails, phones = _known(tid)
    seen_e, seen_p = set(), set()
    out = []
    for i, r in enumerate(rows):
        r = {**r, "line": r.get("line") or i + 2}
        why = _problem(r)
        if why:
            r.update(status="invalid", reason=why)
        elif (r["email"] and r["email"] in emails) or (r["phone"] and r["phone"] in phones):
            r.update(status="existing", reason="Une carte existe déjà pour cet email ou ce téléphone")
        elif (r["email"] and r["email"] in seen_e) or (r["phone"] and r["phone"] in seen_p):
            r.update(status="duplicate", reason="Déjà plus haut dans le fichier")
        else:
            r.update(status="new", reason="")
        if r["status"] == "new":
            if r["email"]:
                seen_e.add(r["email"])
            if r["phone"]:
                seen_p.add(r["phone"])
        out.append(r)
    return out


# ---------------------------------------------------------------------------
# Limits
# ---------------------------------------------------------------------------
def _month(at: datetime) -> datetime:
    return at.replace(day=1, hour=0, minute=0, second=0, microsecond=0)


def invites_used_this_month(tid: str, at: Optional[datetime] = None) -> int:
    at = at or _now()
    try:
        return _db.client_invitations.count_documents(
            {"tenant_id": tid, "status": "sent", "sent_at": {"$gte": _month(at)}})
    except Exception:
        return 0


def limits(tid: str, at: Optional[datetime] = None) -> Dict[str, Any]:
    """What this shop may still do this month: cards, active clients, emails.
    `None` = no limit."""
    import server as core
    at = at or _now()
    t = _db.tenants.find_one({"id": tid}) or {}
    f = core.features_for(t)
    cap = int(f.get("max_customers") or 0)
    monthly = int(f.get("new_cards_per_month") or 0)
    invites = int(f.get("invite_emails_per_month") or 0)
    issued = _db.customers.count_documents({"tenant_id": tid, "created_at": {"$gte": _month(at)}})
    used = invites_used_this_month(tid, at)
    active_left = max(0, cap - core.count_active_customers(tid, at)) if cap else None
    month_left = max(0, monthly - issued) if monthly else None
    cards_left = min(x for x in (active_left, month_left) if x is not None) \
        if active_left is not None or month_left is not None else None
    return {"cards_left": cards_left, "active_left": active_left, "month_cards_left": month_left,
            "invites_limit": invites or None, "invites_used": used,
            "invites_left": max(0, invites - used) if invites else None}


# ---------------------------------------------------------------------------
# The emails
# ---------------------------------------------------------------------------
def _base() -> str:
    import os
    return (os.getenv("PUBLIC_APP_URL") or "https://app.fidclic.com").rstrip("/")


def _email(cust: Dict[str, Any], shop: str, relance: bool) -> Dict[str, Any]:
    first = (cust.get("name") or "").split(" ")[0] or "Bonjour"
    card = f"{_base()}/card/{cust['barcode_id']}?src=invite"
    stop = f"{_base()}/api/public/email/unsubscribe/{cust['email_token']}"
    shop_h, first_h = html.escape(shop), html.escape(first)
    if relance:
        subject = f"Votre carte {shop} vous attend"
        lead = (f"{shop_h} vous a préparé une carte de fidélité, et elle n'attend plus que vous. "
                "Ajoutez-la à votre téléphone : vos points s'y cumulent dès votre prochain passage.")
    else:
        subject = f"{shop} vous offre votre carte de fidélité"
        lead = (f"{shop_h} vous a inscrit(e) à son programme de fidélité. Votre carte est prête : "
                "ajoutez-la à votre téléphone en un clic, et cumulez des avantages à chaque passage.")
    body = f"""<!doctype html><html><body style="margin:0;background:#F5F8FE;font-family:Arial,Helvetica,sans-serif;color:#0B1C3F">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#F5F8FE;padding:28px 12px"><tr><td align="center">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="max-width:520px;background:#ffffff;border:1px solid #E6EEFB;border-radius:16px;padding:28px">
<tr><td style="font-size:13px;font-weight:bold;color:#2F6BFF;letter-spacing:.04em;text-transform:uppercase">{shop_h}</td></tr>
<tr><td style="padding-top:10px;font-size:22px;font-weight:bold">Bonjour {first_h},</td></tr>
<tr><td style="padding-top:12px;font-size:15px;line-height:1.55;color:#334155">{lead}</td></tr>
<tr><td style="padding:24px 0 8px"><a href="{card}" style="display:inline-block;background:#2F6BFF;color:#ffffff;text-decoration:none;font-weight:bold;font-size:15px;padding:13px 22px;border-radius:12px">Ajouter ma carte</a></td></tr>
<tr><td style="font-size:12px;line-height:1.5;color:#94A3B8">Le bouton ne marche pas ? Copiez ce lien : {card}</td></tr>
</table>
<p style="max-width:520px;font-size:11px;line-height:1.5;color:#94A3B8;margin:14px auto 0">
Vous recevez cet email parce que vous êtes client(e) de {shop_h}. Carte de fidélité propulsée par FidClic.<br>
<a href="{stop}" style="color:#64748B">Ne plus recevoir d'emails de {shop_h}</a></p>
</td></tr></table></body></html>"""
    text = (f"Bonjour {first},\n\n{html.unescape(re.sub('<[^>]+>', '', lead))}\n\nAjouter ma carte : {card}\n\n"
            f"Ne plus recevoir d'emails de {shop} : {stop}")
    return {"to": cust["email"], "to_name": cust.get("name") or None, "subject": subject, "html": body, "text": text}


def _send(tid: str, people: List[Dict[str, Any]], kind: str, by: Optional[str],
          import_id: Optional[str] = None) -> Dict[str, Any]:
    """Email these clients (within the month's quota) and log each attempt."""
    from services import mailer
    at = _now()
    lim = limits(tid, at)
    shop = ((_db.tenants.find_one({"id": tid}) or {}).get("campaign_sender_name")
            or (_db.tenants.find_one({"id": tid}) or {}).get("name") or "Votre commerce")
    eligible = [c for c in people if c.get("email") and not c.get("email_opt_out")]
    over = 0
    if lim["invites_left"] is not None and len(eligible) > lim["invites_left"]:
        over = len(eligible) - lim["invites_left"]
        eligible = eligible[:lim["invites_left"]]
    if not mailer.ready():
        return {"sent": 0, "failed": 0, "over_quota": over, "mailer_ready": False}
    for c in eligible:
        if not c.get("email_token"):
            c["email_token"] = secrets.token_urlsafe(18)
            _db.customers.update_one({"id": c["id"]}, {"$set": {"email_token": c["email_token"]}})
    results = mailer.send_batch([_email(c, shop, kind == "relance") for c in eligible], sender_name=shop)
    sent = failed = 0
    for c, r in zip(eligible, results):
        ok = bool(r.get("sent"))
        sent += ok
        failed += not ok
        _db.client_invitations.insert_one({
            "id": str(uuid.uuid4()), "tenant_id": tid, "customer_id": c["id"], "email": c["email"],
            "kind": kind, "status": "sent" if ok else "failed", "reason": r.get("reason"),
            "sent_at": at, "by": by, "import_id": import_id})
    return {"sent": sent, "failed": failed, "over_quota": over, "mailer_ready": True}


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@router.post("/api/owner/clients/import/preview")
async def preview(file: UploadFile = File(...), request: Request = None,
                  token_data: TokenData = Depends(require_role(IMPORTERS, partial_ok=True))):
    raw = await file.read()
    if len(raw) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="Fichier trop lourd (5 Mo au maximum).")
    table = [r for r in _table(file.filename or "", raw) if any(c for c in r)]
    if len(table) < 2:
        raise HTTPException(status_code=400, detail="Le fichier est vide ou n'a qu'une ligne d'en-tête.")
    cols = _columns(table[0])
    if not ({"name", "first", "last"} & set(cols)) or not ({"email", "phone"} & set(cols)):
        raise HTTPException(status_code=400, detail=(
            "Colonnes introuvables. Le fichier doit avoir une colonne Nom (ou Prénom et Nom) "
            "et une colonne Email ou Téléphone."))
    body = table[1:]
    if len(body) > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"{MAX_ROWS} clients au maximum par import.")
    rows = classify(token_data.tenant_id, [{**_row(c, cols), "line": i + 2} for i, c in enumerate(body)])
    counts = {k: sum(1 for r in rows if r["status"] == k) for k in STATUS_LABELS}
    from services import mailer
    return {"columns": {f: table[0][i] for f, i in cols.items()}, "rows": rows, "counts": counts,
            "with_email": sum(1 for r in rows if r["status"] == "new" and r["email"]),
            "limits": limits(token_data.tenant_id), "mailer_ready": mailer.ready(),
            "status_labels": STATUS_LABELS}


class ImportRow(BaseModel):
    name: str = ""
    email: str = ""
    phone: str = ""
    birthday: Optional[str] = None
    postal_code: str = ""
    city: str = ""
    line: Optional[int] = None


class Commit(BaseModel):
    rows: List[ImportRow] = Field(..., max_length=MAX_ROWS)
    consent: bool = False
    invite: bool = False


@router.post("/api/owner/clients/import/commit")
def commit(body: Commit, request: Request = None,
           token_data: TokenData = Depends(require_role(IMPORTERS, partial_ok=True))):
    if not body.consent:
        raise HTTPException(status_code=400, detail=(
            "Confirmez que vos clients ont accepté de rejoindre votre programme de fidélité."))
    import server as core
    import models
    tid = token_data.tenant_id
    at = _now()
    scope = S.resolve(token_data, request)
    # Re-read every line as the server sees it: the browser is never trusted.
    cleaned = [{"name": re.sub(r"\s+", " ", r.name.strip())[:80], "email": r.email.strip().lower()[:120],
                "phone": _phone(r.phone)[:20], "birthday": _mmdd(r.birthday),
                "postal_code": re.sub(r"\s", "", r.postal_code)[:10], "city": r.city.strip()[:80],
                "line": r.line} for r in body.rows]
    rows = classify(tid, cleaned)
    fresh = [r for r in rows if r["status"] == "new"]
    lim = limits(tid, at)
    over_cards = 0
    if lim["cards_left"] is not None and len(fresh) > lim["cards_left"]:
        over_cards = len(fresh) - lim["cards_left"]
        fresh = fresh[:lim["cards_left"]]

    import_id = str(uuid.uuid4())
    # The establishment the clients belong to: the one chosen in the sidebar,
    # or for an administrateur of several, the first of theirs (else they
    # would import clients they then could not see).
    store = scope.single or (None if scope.is_all else (list(scope.ids or []) or [None])[0])
    created = []
    for r in fresh:
        city = r["city"]
        if not city and r["postal_code"]:
            try:
                from services import communes
                label = communes.label_for(r["postal_code"])
                city = label if label and label != r["postal_code"] else ""
            except Exception:
                city = ""
        doc = models.Customer(
            id=str(uuid.uuid4()), tenant_id=tid, barcode_id="FT-" + uuid.uuid4().hex[:8].upper(),
            name=r["name"], email=r["email"] or "", phone=r["phone"] or "", postal_code=r["postal_code"] or "",
            birthday=r["birthday"] or "", acquisition_source="import", branch_id=store,
        ).model_dump()
        doc.update({
            "created_at": at, "signup_channel": "import", "import_id": import_id,
            "email_token": secrets.token_urlsafe(18),
            # The merchant's word that this client agreed — the card's consent record.
            "consent": {"accepted": True, "ts": at, "source": "import", "by": token_data.email,
                        "policy_version": core.CURRENT_PRIVACY_POLICY_VERSION},
        })
        if city:
            doc["city"] = city
        _db.customers.insert_one(dict(doc))
        created.append(doc)
    core.log_audit("clients.import", actor=token_data.email, actor_type="user", tenant_id=tid,
                   target=import_id, request=request,
                   meta={"created": len(created), "lines": len(rows), "invite": body.invite})

    mail = {"sent": 0, "failed": 0, "over_quota": 0, "mailer_ready": None}
    if body.invite and created:
        mail = _send(tid, created, "invite", token_data.email, import_id)
    counts = {k: sum(1 for r in rows if r["status"] == k) for k in STATUS_LABELS}
    return {"created": len(created), "over_limit": over_cards, "counts": counts,
            "emails": mail, "limits": limits(tid), "import_id": import_id}


def _invite_rows(tid: str, mine: Optional[set]) -> Dict[str, Dict[str, Any]]:
    """Per invited client: first invitation, relance sent or not."""
    out: Dict[str, Dict[str, Any]] = {}
    for inv in _db.client_invitations.find({"tenant_id": tid, "status": "sent"},
                                           {"_id": 0, "customer_id": 1, "kind": 1, "sent_at": 1}):
        cid = inv["customer_id"]
        if mine is not None and cid not in mine:
            continue
        row = out.setdefault(cid, {"invited_at": None, "relanced": False})
        t = _aware(inv.get("sent_at"))
        if inv.get("kind") == "relance":
            row["relanced"] = True
        elif t and (row["invited_at"] is None or t < row["invited_at"]):
            row["invited_at"] = t
    return out


def _opened(tid: str, ids: List[str]) -> set:
    """Clients who opened their card, or came in, since being invited."""
    seen = set()
    for o in _db.card_opens.find({"tenant_id": tid, "customer_id": {"$in": ids}}, {"_id": 0, "customer_id": 1}):
        seen.add(o["customer_id"])
    for c in _db.customers.find({"id": {"$in": ids}, "visits": {"$gt": 0}}, {"_id": 0, "id": 1}):
        seen.add(c["id"])
    return seen


def relance_candidates(tid: str, mine: Optional[set] = None, at: Optional[datetime] = None) -> List[Dict[str, Any]]:
    at = at or _now()
    inv = _invite_rows(tid, mine)
    due = [cid for cid, r in inv.items()
           if r["invited_at"] and not r["relanced"] and at - r["invited_at"] >= timedelta(days=RELANCE_AFTER_DAYS)]
    opened = _opened(tid, due)
    people = {c["id"]: c for c in _db.customers.find(
        {"tenant_id": tid, "id": {"$in": [d for d in due if d not in opened]}},
        {"_id": 0, "id": 1, "name": 1, "email": 1, "email_opt_out": 1, "deleted_at": 1})}
    return sorted(({"id": cid, "name": c.get("name") or "Client", "email": c.get("email"),
                    "invited_at": inv[cid]["invited_at"].isoformat()}
                   for cid, c in people.items()
                   if c.get("email") and not c.get("email_opt_out") and not c.get("deleted_at")),
                  key=lambda r: r["invited_at"])


@router.get("/api/owner/clients/invitations")
def invitations(request: Request = None, token_data: TokenData = Depends(require_role(READERS, partial_ok=True))):
    from services import mailer
    tid = token_data.tenant_id
    mine = S.customer_ids(S.resolve(token_data, request))
    inv = _invite_rows(tid, mine)
    opened = _opened(tid, list(inv))
    return {"invited": len(inv), "opened": len(opened), "waiting": len(inv) - len(opened),
            "relance_ready": len(relance_candidates(tid, mine)), "relance_after_days": RELANCE_AFTER_DAYS,
            "limits": limits(tid), "mailer_ready": mailer.ready()}


@router.get("/api/owner/clients/invitations/relance")
def relance_list(request: Request = None, token_data: TokenData = Depends(require_role(READERS, partial_ok=True))):
    tid = token_data.tenant_id
    rows = relance_candidates(tid, S.customer_ids(S.resolve(token_data, request)))
    return {"rows": rows, "limits": limits(tid), "relance_after_days": RELANCE_AFTER_DAYS}


class Relance(BaseModel):
    customer_ids: List[str] = Field(..., min_length=1, max_length=MAX_ROWS)


@router.post("/api/owner/clients/invitations/relance")
def relance_send(body: Relance, request: Request = None,
                 token_data: TokenData = Depends(require_role(IMPORTERS, partial_ok=True))):
    """Send the relance to the clients the merchant ticked — and only to those
    who still qualify (a client who opened their card meanwhile is skipped)."""
    tid = token_data.tenant_id
    allowed = {r["id"] for r in relance_candidates(tid, S.customer_ids(S.resolve(token_data, request)))}
    ids = [i for i in body.customer_ids if i in allowed]
    if not ids:
        raise HTTPException(status_code=400, detail="Aucun de ces clients n'attend de relance.")
    people = list(_db.customers.find({"tenant_id": tid, "id": {"$in": ids}}, {"_id": 0}))
    return {**_send(tid, people, "relance", token_data.email), "limits": limits(tid)}


@router.get("/api/public/email/unsubscribe/{token}", response_class=HTMLResponse)
def unsubscribe(token: str):
    cust = _db.customers.find_one({"email_token": token}) if token and len(token) >= 16 else None
    if cust:
        _db.customers.update_one({"id": cust["id"]}, {"$set": {"email_opt_out": True, "email_opt_out_at": _now()}})
        shop = (_db.tenants.find_one({"id": cust.get("tenant_id")}) or {}).get("name") or "ce commerce"
        msg = f"C'est noté : vous ne recevrez plus d'emails de {html.escape(shop)}. Votre carte de fidélité reste valable."
    else:
        msg = "Ce lien n'est plus valable."
    return HTMLResponse(f"""<!doctype html><html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Désinscription</title></head><body style="margin:0;background:#F5F8FE;font-family:Arial,Helvetica,sans-serif;color:#0B1C3F">
<div style="max-width:440px;margin:15vh auto;background:#fff;border:1px solid #E6EEFB;border-radius:16px;padding:28px;text-align:center">
<p style="font-size:16px;line-height:1.5">{msg}</p></div></body></html>""")
